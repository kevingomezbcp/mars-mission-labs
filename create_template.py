import pptx

def create_clean_template(src_pptx, template_pptx):
    prs = pptx.Presentation(src_pptx)
    
    # Drop all slides and their part relationships
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.rId
        # Drop relationship from presentation part
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)
        
    prs.save(template_pptx)
    print(f"Saved clean template with {len(prs.slides)} slides to {template_pptx}")

create_clean_template("Miercoles de IA - Agent Framework - 10-06-2026.pptx", "template_miercoles_ia.pptx")
