import sys
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation("RAG_Microsoft_Agent_Framework_Transit.pptx")
print(f"Total Slides: {len(prs.slides)}")
assert len(prs.slides) == 14, "Expected 14 slides!"

for idx, s in enumerate(prs.slides):
    print(f"\n{'='*25} SLIDE {idx+1}: {s.slide_layout.name} {'='*25}")
    for sh_idx, sh in enumerate(s.shapes):
        t_preview = ""
        if sh.has_text_frame:
            texts = [p.text.strip().replace('\n', ' ') for p in sh.text_frame.paragraphs if p.text.strip()]
            t_preview = " | ".join(texts)[:80]
        elif sh.has_table:
            t_preview = f"[Table {len(sh.table.rows)}x{len(sh.table.columns)}]"
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            t_preview = "[Picture]"
        print(f"  Shape {sh_idx}: '{sh.name}' ({sh.shape_type}) -> {t_preview}")

print("\n" + "="*50)
print("VERIFICATION COMPLETED: ALL 14 SLIDES VALIDATED!")
print("="*50)
